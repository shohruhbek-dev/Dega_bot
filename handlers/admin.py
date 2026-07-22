import io
import asyncio
import logging
from datetime import datetime
from aiogram import Router, F, Bot
from aiogram.filters import Command
from aiogram.types import Message, CallbackQuery, BufferedInputFile

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

import database as db
from config import ADMIN_IDS

logger = logging.getLogger(__name__)
router = Router()


def is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS


def _thin_border() -> Border:
    side = Side(style="thin", color="CCCCCC")
    return Border(left=side, right=side, top=side, bottom=side)


def _build_excel(rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Buyurtmalar"

    # Sarlavha uslubi
    header_font  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill  = PatternFill("solid", fgColor="1B3A6B")   # Navy Blue
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Ma'lumot uslubi
    data_font    = Font(name="Calibri", size=10)
    data_align   = Alignment(horizontal="left", vertical="center", wrap_text=False)
    center_align = Alignment(horizontal="center", vertical="center")

    HEADERS = [
        "№", "User ID", "Username", "Ism-Familiya",
        "Telefon", "Latitude", "Longitude",
        "Xizmat turi", "Ro'yxatdan o'tgan", "Buyurtma vaqti",
    ]

    ws.row_dimensions[1].height = 24

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font   = header_font
        cell.fill   = header_fill
        cell.alignment = header_align
        cell.border = _thin_border()

    # Ma'lumotlarni yozish
    for row_idx, row in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 18
        values = [
            row_idx - 1,
            row.get("user_id", ""),
            ("@" + row["username"]) if row.get("username") else "—",
            row.get("full_name") or "—",
            row.get("phone") or "—",
            row.get("latitude") or "—",
            row.get("longitude") or "—",
            row.get("service_name") or "—",
            "✅" if row.get("registered") else "❌",
            (row.get("created_at") or "")[:16],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font   = data_font
            cell.border = _thin_border()
            cell.alignment = (
                center_align
                if col_idx in (1, 6, 7, 9)
                else data_align
            )

    # Ustun kengliklari avtomatik moslashtirish
    col_min_widths = [5, 14, 18, 24, 16, 12, 12, 30, 14, 18]
    for col_idx, min_w in enumerate(col_min_widths, start=1):
        letter = get_column_letter(col_idx)
        max_w = min_w
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_w = max(max_w, len(str(cell.value)) + 2)
        ws.column_dimensions[letter].width = min(max_w, 40)

    # Muzlatish (birinchi qator)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def _send_excel(
    bot: Bot,
    chat_id: int,
) -> None:
    rows = await db.get_orders_with_users()
    if not rows:
        await bot.send_message(chat_id, "Bazada buyurtma yo'q.")
        return

    data = _build_excel(rows)
    filename = f"DEGA_orders_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    file = BufferedInputFile(data, filename=filename)

    await bot.send_document(
        chat_id,
        file,
        caption=(
            f"<b>DEGA — Buyurtmalar hisoboti</b>\n"
            f"Jami buyurtmalar: <b>{len(rows)}</b>\n"
            f"Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ),
        parse_mode="HTML",
    )


# ── /export ──────────────────────────────────────────────────────────────────

@router.message(Command("export"))
async def export_cmd(message: Message, bot: Bot) -> None:
    if not is_admin(message.from_user.id):
        await message.answer("Ruxsat yo'q.")
        return
    await message.answer("Tayyorlanmoqda...")
    await _send_excel(bot, message.chat.id)


@router.callback_query(F.data == "admin_export")
async def export_callback(callback: CallbackQuery, bot: Bot) -> None:
    if not is_admin(callback.from_user.id):
        await callback.answer("Ruxsat yo'q.", show_alert=True)
        return
    await callback.answer("Tayyorlanmoqda...")
    await _send_excel(bot, callback.message.chat.id)


# ── /send ─────────────────────────────────────────────────────────────────────

@router.message(Command("send"))
async def broadcast_cmd(message: Message, bot: Bot) -> None:
    if not is_admin(message.from_user.id):
        await message.answer("Ruxsat yo'q.")
        return

    parts = message.text.split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await message.answer(
            "Foydalanish: <code>/send Xabar matni</code>",
            parse_mode="HTML",
        )
        return

    broadcast_text = parts[1].strip()
    user_ids = await db.get_registered_user_ids()

    if not user_ids:
        await message.answer("Ro'yxatdan o'tgan foydalanuvchi yo'q.")
        return

    status = await message.answer(f"Yuborilmoqda... (0/{len(user_ids)})")
    success, failed = 0, 0

    for i, uid in enumerate(user_ids):
        try:
            await bot.send_message(uid, broadcast_text, parse_mode="HTML")
            success += 1
        except Exception as e:
            logger.warning(f"Broadcast xato uid={uid}: {e}")
            failed += 1

        if (i + 1) % 25 == 0:
            try:
                await status.edit_text(f"Yuborilmoqda... ({i + 1}/{len(user_ids)})")
            except Exception:
                pass

        await asyncio.sleep(0.05)

    await status.edit_text(
        f"<b>Broadcast yakunlandi</b>\n\n"
        f"Jami: {len(user_ids)}\n"
        f"Muvaffaqiyatli: {success}\n"
        f"Xato: {failed}",
        parse_mode="HTML",
    )


# ── /stats ────────────────────────────────────────────────────────────────────

@router.message(Command("stats"))
async def stats_cmd(message: Message) -> None:
    if not is_admin(message.from_user.id):
        await message.answer("Ruxsat yo'q.")
        return

    all_users  = await db.get_all_users()
    registered = [u for u in all_users if u.get("registered")]
    orders     = await db.get_orders_with_users()

    lang_counts: dict[str, int] = {}
    for u in registered:
        lg = u.get("language", "uz")
        lang_counts[lg] = lang_counts.get(lg, 0) + 1

    lang_lines = "\n".join(
        f"  {lg.upper()}: {cnt}" for lg, cnt in lang_counts.items()
    )

    await message.answer(
        f"<b>DEGA Bot — Statistika</b>\n\n"
        f"Jami foydalanuvchilar: {len(all_users)}\n"
        f"Ro'yxatdan o'tganlar: {len(registered)}\n"
        f"Tugallamagan: {len(all_users) - len(registered)}\n"
        f"Jami buyurtmalar: {len(orders)}\n\n"
        f"Tillar bo'yicha:\n{lang_lines or '  —'}\n\n"
        f"Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        parse_mode="HTML",
    )
